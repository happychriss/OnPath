#!/usr/bin/env python3
"""
OnPath texts EN + DE: review sheet round trip (docs/i18n.md).

  i18n.py export [--out DIR]          string files + tools/i18n/meta → review.csv + review.xlsx
                                      (default DIR: external-docs/translation, private, not in git)
  i18n.py import FILE [--dry-run]     owner columns "EN neu (Owner)" / "DE neu (Owner)" of a CSV / xlsx → string files;
                                      checks the placeholders, prints a diff
  i18n.py check                       placeholders, German texts present, metadata present (like StringFilesTest)

Texts: strings/src/commonMain/composeResources/values/strings_<area>.xml (English) and values-de/ (German).
Rows: one per string; a plural gives one row per form, key "name#one" / "name#other".
Needs openpyxl for .xlsx (pip install openpyxl); CSV works without it.
"""
import argparse
import csv
import difflib
import json
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
RES = ROOT / "strings/src/commonMain/composeResources"
META = ROOT / "tools/i18n/meta"
GLOSSARY = ROOT / "tools/i18n/glossary.tsv"
DEFAULT_OUT = ROOT / "external-docs/translation"

COLUMNS = ["Key", "Priorität", "Sicherheit", "Bereich", "Kontext", "English", "Deutsch (Vorschlag)",
           "EN neu (Owner)", "DE neu (Owner)", "Kommentar (Owner)", "Status"]
PLACEHOLDER = re.compile(r"%(\d)\$([ds])")

BACKGROUND = [
    ("Hintergrund / Background", ""),
    ("EN",
     "OnPath is a simple, fully offline hiking and biking app for Android: plan a route on the map, see time, "
     "distance and difficulty, follow it with navigation and record the tour – like the outdoor planning apps many "
     "people know, but much simpler and without an account. Users: hikers and cyclists, also in German-speaking "
     "countries (Germany, Austria, South Tyrol, Switzerland). Tone: friendly and factual, short, no marketing words, "
     "no exclamation marks. UI: a phone screen – buttons and chips are short (see 'Kontext' for limits). "
     "Placeholders such as %1$s or %2$d must stay (they are replaced by names, distances, counts); rows named "
     "'key#one' / 'key#other' are the singular / plural forms of the same text."),
    ("DE",
     "OnPath ist eine einfache Wander- und Rad-App für Android, die komplett offline funktioniert: Route auf der "
     "Karte planen, Zeit, Strecke und Schwierigkeit sehen, mit Navigation folgen und die Tour aufzeichnen – wie die "
     "bekannten Outdoor-Planungs-Apps, nur deutlich einfacher und ohne Konto. Nutzer: Wanderer und Radfahrer, auch im "
     "deutschsprachigen Raum (Deutschland, Österreich, Südtirol, Schweiz). Ton: freundlich mit „du“, sachlich, kurz, "
     "keine Werbesprache, keine Ausrufezeichen. Oberfläche: Handy-Bildschirm – Buttons und Chips sind kurz (Grenzen "
     "unter 'Kontext'). Platzhalter wie %1$s oder %2$d müssen bleiben (sie werden durch Namen, Strecken, Zahlen "
     "ersetzt); Zeilen 'key#one' / 'key#other' sind Einzahl / Mehrzahl desselben Textes."),
    ("Wie prüfen / How to review",
     "Spalte 'EN neu' oder 'DE neu' nur ausfüllen, wenn der Text anders sein soll; 'Kommentar' für Fragen; 'Status' "
     "z. B. ok / geändert / offen. Sicherheit: sicher / bitte prüfen / Zweifel (Grund im Kontext). Priorität 1 = fast "
     "bei jeder Nutzung zu sehen (Planen, Aktivität, Navigation), 2 = regelmäßig (Einstellungen, Meine Routen, "
     "Dialoge), 3 = selten (Fehler, Sonderfälle)."),
]


# ---------------------------------------------------------------- reading

def read_values(dirname):
    """{file name: [(key, kind, form, text)]} in file order; kind = string | plural, form = '' or quantity."""
    out = {}
    d = RES / dirname
    for f in sorted(d.glob("*.xml")):
        rows = []
        for e in ET.parse(f).getroot():
            if not isinstance(e.tag, str):
                continue
            name = e.get("name")
            if e.tag == "string":
                rows.append((name, "string", "", e.text or ""))
            elif e.tag == "plurals":
                for item in e.findall("item"):
                    rows.append((name, "plural", item.get("quantity"), item.text or ""))
        out[f.name] = rows
    return out


def read_meta():
    meta = {}
    for f in sorted(META.glob("*.json")):
        for k, v in json.loads(f.read_text(encoding="utf-8")).items():
            meta[k] = v
    return meta


def row_key(key, form):
    return f"{key}#{form}" if form else key


def collect():
    en, de, meta = read_values("values"), read_values("values-de"), read_meta()
    de_map = {}
    for rows in de.values():
        for key, kind, form, text in rows:
            de_map[row_key(key, form)] = text
    rows = []
    for fname, items in en.items():
        for order, (key, kind, form, text) in enumerate(items):
            m = meta.get(key, {})
            ctx = m.get("context", "")
            if m.get("max"):
                ctx += f" (max. ~{m['max']} Zeichen)"
            if m.get("doubt"):
                ctx += f" – Zweifel: {m['doubt']}"
            rows.append({
                "Key": row_key(key, form), "Priorität": m.get("priority", 3),
                "Sicherheit": m.get("confidence", "bitte prüfen"), "Bereich": m.get("area", fname[8:-4]),
                "Kontext": ctx.strip(), "English": text, "Deutsch (Vorschlag)": de_map.get(row_key(key, form), ""),
                "EN neu (Owner)": "", "DE neu (Owner)": "", "Kommentar (Owner)": "", "Status": "",
                "_file": fname, "_order": order,
            })
    rows.sort(key=lambda r: (int(r["Priorität"]), r["Bereich"].lower(), r["_file"], r["_order"]))
    return rows, meta


def glossary():
    lines = GLOSSARY.read_text(encoding="utf-8").splitlines()
    return [l.split("\t") for l in lines if l.strip()]


# ---------------------------------------------------------------- export

def export(out_dir: Path):
    rows, _ = collect()
    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "review.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("openpyxl missing: wrote only review.csv")
        return rows
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Texte"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r[c] for c in COLUMNS])
    widths = [30, 9, 13, 18, 50, 42, 42, 30, 30, 30, 12]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wdt
    head = PatternFill("solid", fgColor="2E6B3F")
    owner = PatternFill("solid", fgColor="FFF6D5")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"  # a text, not a formula
        for c in row[7:11]:
            c.fill = owner
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    bg = wb.create_sheet("Hintergrund", 0)
    for title, text in BACKGROUND:
        bg.append([title, text])
    bg.column_dimensions["A"].width = 26
    bg.column_dimensions["B"].width = 120
    for row in bg.iter_rows():
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    bg["A1"].font = Font(bold=True, size=14)

    gl = wb.create_sheet("Glossar", 1)
    for line in glossary():
        gl.append(line + [""] * (3 - len(line)))
    for col, wdt in zip("ABC", (30, 30, 80)):
        gl.column_dimensions[col].width = wdt
    for c in gl[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
    gl.freeze_panes = "A2"
    wb.save(out_dir / "review.xlsx")
    return rows


# ---------------------------------------------------------------- import

def read_sheet(path: Path):
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["Texte"] if "Texte" in wb.sheetnames else wb.active
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip() for h in next(it)]
        return [{h: ("" if v is None else str(v)) for h, v in zip(head, row)} for row in it]
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def esc(text):
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def replace_in_file(src: str, key: str, form: str, new: str):
    """Replaces one text in the raw XML (keeps the file's layout); returns the new source or None."""
    if not form:
        pat = re.compile(r'(<string name="%s">)(.*?)(</string>)' % re.escape(key), re.S)
        m = pat.search(src)
        return None if not m else src[:m.start(2)] + esc(new) + src[m.end(2):]
    block = re.compile(r'<plurals name="%s">.*?</plurals>' % re.escape(key), re.S).search(src)
    if not block:
        return None
    item = re.compile(r'(<item quantity="%s">)(.*?)(</item>)' % re.escape(form), re.S).search(src, block.start(), block.end())
    return None if not item else src[:item.start(2)] + esc(new) + src[item.end(2):]


def holders(s):
    return sorted(set(m.group(0) for m in PLACEHOLDER.finditer(s)))


def do_import(path: Path, dry: bool):
    rows = read_sheet(path)
    en, de = read_values("values"), read_values("values-de")
    where = {}
    for lang, files in (("en", en), ("de", de)):
        for fname, items in files.items():
            for key, kind, form, text in items:
                where[(lang, row_key(key, form))] = (fname, key, form, text)
    changes, errors = [], []
    for r in rows:
        k = (r.get("Key") or "").strip()
        if not k:
            continue
        new_en, new_de = (r.get("EN neu (Owner)") or "").strip(), (r.get("DE neu (Owner)") or "").strip()
        if not new_en and not new_de:
            continue
        cur_en = where.get(("en", k))
        if not cur_en:
            errors.append(f"{k}: unknown key")
            continue
        cur_de = where.get(("de", k))
        en_text = new_en or cur_en[3]
        de_text = new_de or (cur_de[3] if cur_de else "")
        is_plural_one = k.endswith("#one")
        if de_text and not is_plural_one and holders(en_text) != holders(de_text):
            errors.append(f"{k}: placeholders differ – EN {holders(en_text)} vs DE {holders(de_text)}")
            continue
        if new_en and new_en != cur_en[3]:
            changes.append(("values", cur_en[0], cur_en[1], cur_en[2], new_en))
        if new_de and (not cur_de or new_de != cur_de[3]):
            if not cur_de:
                errors.append(f"{k}: no German entry to replace")
                continue
            changes.append(("values-de", cur_de[0], cur_de[1], cur_de[2], new_de))
    files = {}
    for d, fname, key, form, new in changes:
        p = RES / d / fname
        src = files.get(p) or p.read_text(encoding="utf-8")
        out = replace_in_file(src, key, form, new)
        if out is None:
            errors.append(f"{row_key(key, form)}: not found in {d}/{fname}")
            continue
        files[p] = out
    for p, new_src in files.items():
        old = p.read_text(encoding="utf-8")
        sys.stdout.writelines(difflib.unified_diff(old.splitlines(True), new_src.splitlines(True),
                                                   str(p.relative_to(ROOT)), str(p.relative_to(ROOT)) + " (new)"))
        if not dry:
            p.write_text(new_src, encoding="utf-8")
    for e in errors:
        print("ERROR", e, file=sys.stderr)
    print(f"{len(changes)} change(s){' (dry run)' if dry else ''}, {len(errors)} error(s)")
    return 1 if errors else 0


# ---------------------------------------------------------------- check

def check():
    rows, meta = collect()
    bad = []
    for r in rows:
        k = r["Key"]
        if not r["Deutsch (Vorschlag)"]:
            bad.append(f"{k}: no German text")
        elif not k.endswith("#one") and holders(r["English"]) != holders(r["Deutsch (Vorschlag)"]):
            bad.append(f"{k}: placeholders differ")
        if k.split("#")[0] not in meta:
            bad.append(f"{k}: no metadata in tools/i18n/meta")
    for b in bad:
        print(b)
    print(f"{len(rows)} rows, {len(bad)} problem(s)")
    return 1 if bad else 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("--out", type=Path, default=DEFAULT_OUT)
    i = sub.add_parser("import")
    i.add_argument("file", type=Path)
    i.add_argument("--dry-run", action="store_true")
    sub.add_parser("check")
    a = ap.parse_args()
    if a.cmd == "export":
        rows = export(a.out)
        print(f"{len(rows)} rows → {a.out}/review.csv, review.xlsx")
        return 0
    if a.cmd == "import":
        return do_import(a.file, a.dry_run)
    return check()


if __name__ == "__main__":
    sys.exit(main())
